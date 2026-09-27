"""報告生成：可疑交易分析報告草稿 + 稽核工作底稿（Excel）。

產生方式（config.yaml → report.llm.provider）
- template：以規則證據、SHAP 解釋與交易統計套入固定模板，完全離線、結果可重現。
- anthropic / openai：把同一份「結構化事實」交給大型語言模型改寫成流暢的敘述；
  模型只能使用提供的事實，不得自行杜撰。
- auto（預設）：偵測到 ANTHROPIC_API_KEY 或 OPENAI_API_KEY 就用 LLM，否則用模板；
  LLM 呼叫失敗時自動退回模板，確保比賽現場沒有網路也能展示。
"""
from __future__ import annotations

import io
import json
import os
import re
from datetime import datetime

import numpy as np
import pandas as pd
import requests

from .features import FEATURE_NAMES_ZH
from .data_loader import CASH
from .rules import RULE_DESCRIPTIONS, RULE_REFS, RULES

DISCLAIMER = (
    "本報告由系統依交易資料自動產生，屬稽核輔助資訊，須經稽核／法遵人員覆核確認後，"
    "始得作為內部處置或向主管機關申報之依據。"
)


def risk_level(p: float) -> str:
    if p >= 0.8:
        return "高"
    if p >= 0.5:
        return "中"
    return "低"


def _fmt(x: float) -> str:
    return f"{x:,.0f}"


# ---------------------------------------------------------------------------
# 1. 整理結構化事實（模板與 LLM 共用）
# ---------------------------------------------------------------------------
def collect_facts(account_id: str, results: pd.DataFrame, shap_df: pd.DataFrame, X: pd.DataFrame,
                  tx: pd.DataFrame, currency_label: str = "元", top_n: int = 5) -> dict:
    r = results.loc[account_id]
    t_in = tx[tx["dst"] == account_id]
    t_out = tx[tx["src"] == account_id]
    t_all = pd.concat([t_in, t_out])

    # SHAP：把風險往上推最多的特徵
    s = shap_df.loc[account_id].drop("_bias", errors="ignore")
    med = X.median()
    pct = X.rank(pct=True).loc[account_id]
    reasons = []
    for feat, val in s.sort_values(ascending=False).head(top_n).items():
        if val <= 0:
            break
        reasons.append(
            {
                "feature": feat,
                "name_zh": FEATURE_NAMES_ZH.get(feat, feat),
                "value": float(X.at[account_id, feat]),
                "population_median": float(med[feat]),
                "percentile": float(pct[feat]),
                "shap": float(val),
            }
        )

    rule_hits = []
    for code, name in RULES.items():
        if bool(r.get(f"{code}_hit", False)):
            rule_hits.append({"code": code, "name": name, "description": RULE_DESCRIPTIONS[code],
                              "reference": RULE_REFS[code], "evidence": str(r.get(f"{code}_evidence", ""))})

    def top_cp(df, col):
        if df.empty:
            return []
        g = df.groupby(col)["amount"].agg(["count", "sum"]).sort_values("sum", ascending=False).head(5)
        out = []
        for cp, row in g.iterrows():
            out.append({
                "account": "（現金）" if cp == CASH else cp,
                "n_tx": int(row["count"]),
                "amount": float(row["sum"]),
                "risk_score": float(results.at[cp, "risk_score"]) if cp in results.index else None,
            })
        return out

    return {
        "account_id": account_id,
        "risk_score": float(r["risk_score"]),
        "risk_level": risk_level(float(r["risk_score"])),
        "risk_rank": int(r["risk_rank"]),
        "n_accounts": int(len(results)),
        "period": f"{t_all['date'].min():%Y-%m-%d} ~ {t_all['date'].max():%Y-%m-%d}" if len(t_all) else "無交易",
        "currency_label": currency_label,
        "summary": {
            "n_in": int(len(t_in)), "amt_in": float(t_in["amount"].sum()),
            "n_out": int(len(t_out)), "amt_out": float(t_out["amount"].sum()),
            "n_src": int(t_in.loc[t_in["src"] != CASH, "src"].nunique()),
            "n_dst": int(t_out.loc[t_out["dst"] != CASH, "dst"].nunique()),
            "active_days": int(t_all["step"].nunique()),
            "cash_in": float(t_in.loc[t_in["src"] == CASH, "amount"].sum()),
            "cash_out": float(t_out.loc[t_out["dst"] == CASH, "amount"].sum()),
        },
        "profile": {k: (str(r[k].date()) if hasattr(r[k], "date") and pd.notna(r[k]) else (None if pd.isna(r[k]) else r[k]))
                    for k in ("customer_type", "open_date", "dormant_days_before") if k in r.index},
        "rule_hits": rule_hits,
        "model_reasons": reasons,
        "top_sources": top_cp(t_in, "src"),
        "top_destinations": top_cp(t_out, "dst"),
        "group_id": int(r.get("group_id", -1)),
        "group_size": int(r.get("group_size", 0)),
        "group_high_risk": int(r.get("group_high_risk", 0)),
    }


# ---------------------------------------------------------------------------
# 2. 模板產生
# ---------------------------------------------------------------------------
def _recommendations(f: dict) -> list[str]:
    lv = f["risk_level"]
    rec = []
    if lv == "高":
        rec += [
            "將本帳戶列為高風險客戶，立即啟動加強客戶審查（EDD），確認開戶目的、資金來源與實際使用人。",
            "調閱開戶文件、登入 IP／裝置紀錄與臨櫃影像，確認是否有帳戶出租、出借或遭收購情形。",
            "依機構內部程序評估是否向法務部調查局申報疑似洗錢交易，並評估暫停網路銀行或非約定轉帳功能。",
        ]
    elif lv == "中":
        rec += [
            "列入持續監控名單，於下一個稽核週期再次評估。",
            "由理專或分行主管聯繫客戶，了解近期交易目的並留存紀錄。",
        ]
    else:
        rec += ["目前風險較低，維持一般監控即可。"]
    names = {h["code"] for h in f["rule_hits"]}
    if "R4" in names:
        rec.append("資金循環路徑上的所有帳戶宜一併調查，確認是否為同一集團控制。")
    if "R2" in names:
        rec.append("清查匯入來源帳戶是否有詐騙報案紀錄，以協助被害人追回款項。")
    if f.get("group_size", 0) >= 3:
        rec.append(f"本帳戶屬於疑似集團 G{f['group_id']}（共 {f['group_size']} 個可疑帳戶，其中另有 {f['group_high_risk']} 個高風險），"
                   "建議以集團為單位擴大查核，並比對是否有共用開戶資料、裝置或 IP。")
    return rec


def render_template(f: dict) -> str:
    cur = f["currency_label"]
    s = f["summary"]
    L = []
    L.append(f"# 疑似人頭帳戶／洗錢交易分析報告（草稿）\n")
    L.append(f"**報告產生時間：** {datetime.now():%Y-%m-%d %H:%M}　**產生方式：** 規則引擎＋AI 模型（模板）\n")
    L.append("## 一、基本資訊\n")
    L.append("| 項目 | 內容 |\n|---|---|")
    L.append(f"| 帳戶 | {f['account_id']} |")
    L.append(f"| AI 風險分數 | {f['risk_score']:.3f}（風險等級：**{f['risk_level']}**） |")
    L.append(f"| 全體排名 | 第 {f['risk_rank']:,} 名／共 {f['n_accounts']:,} 個帳戶 |")
    L.append(f"| 分析期間 | {f['period']} |")
    L.append(f"| 命中紅旗規則 | {len(f['rule_hits'])} 項 |")
    L.append(f"| 所屬疑似集團 | {'G' + str(f['group_id']) + '（' + str(f['group_size']) + ' 個帳戶）' if f['group_id'] > 0 else '無'} |\n")

    L.append("## 二、交易概況\n")
    L.append(
        f"分析期間內，本帳戶共匯入 {s['n_in']} 筆、金額 {_fmt(s['amt_in'])} {cur}（來自 {s['n_src']} 個不同帳戶）；"
        f"匯出 {s['n_out']} 筆、金額 {_fmt(s['amt_out'])} {cur}（匯往 {s['n_dst']} 個不同帳戶），"
        f"共有 {s['active_days']} 天發生交易。"
    )
    if s.get("cash_out") or s.get("cash_in"):
        L.append(f"其中現金存入 {_fmt(s.get('cash_in', 0))} {cur}、現金提領 {_fmt(s.get('cash_out', 0))} {cur}。")
    prof = f.get("profile") or {}
    if prof.get("open_date"):
        L.append(f"帳戶類型：{prof.get('customer_type', '-')}；開戶日：{prof['open_date']}；"
                 f"觀察期前已未往來 {prof.get('dormant_days_before', 0)} 天。")
    if s["amt_in"] > 0:
        L.append(f"匯出金額相當於匯入金額的 {s['amt_out'] / s['amt_in']:.0%}。\n")
    else:
        L.append("")

    L.append("## 三、命中之紅旗指標\n")
    if f["rule_hits"]:
        for i, h in enumerate(f["rule_hits"], 1):
            L.append(f"{i}. **{h['code']} {h['name']}**（{h.get('reference', '')}）：{h['description']}")
            L.append(f"   - 證據：{h['evidence']}")
    else:
        L.append("未命中任何固定規則；本帳戶係由 AI 模型依整體行為型態判定為可疑，屬規則難以涵蓋之新型態，建議優先覆核。")
    L.append("")

    L.append("## 四、AI 模型判斷依據\n")
    L.append("下列為將本帳戶風險分數往上推升最多的因素（TreeSHAP 可解釋性分析）：\n")
    L.append("| 因素 | 本帳戶數值 | 全體中位數 | 與全體帳戶比較 |\n|---|---|---|---|")
    for r in f["model_reasons"]:
        p = r["percentile"]
        cmp_txt = f"高於 {p:.1%} 的帳戶" if p >= 0.5 else f"低於 {1 - p:.1%} 的帳戶"
        L.append(f"| {r['name_zh']} | {r['value']:,.2f} | {r['population_median']:,.2f} | {cmp_txt} |")
    L.append("")

    L.append("## 五、主要資金往來對象\n")
    for title, key in (("主要匯入來源", "top_sources"), ("主要匯出對象", "top_destinations")):
        L.append(f"**{title}**\n")
        if f[key]:
            L.append("| 帳戶 | 筆數 | 金額 | 對方 AI 風險分數 |\n|---|---|---|---|")
            for c in f[key]:
                rs = f"{c['risk_score']:.3f}" if c["risk_score"] is not None else "-"
                L.append(f"| {c['account']} | {c['n_tx']} | {_fmt(c['amount'])} | {rs} |")
        else:
            L.append("（無）")
        L.append("")

    L.append("## 六、稽核建議\n")
    for i, r in enumerate(_recommendations(f), 1):
        L.append(f"{i}. {r}")
    L.append("")
    L.append("## 七、聲明\n")
    L.append(DISCLAIMER)
    return "\n".join(L)


# ---------------------------------------------------------------------------
# 3. LLM 產生（可選）
# ---------------------------------------------------------------------------
LLM_SYSTEM = (
    "你是銀行內部稽核與洗錢防制（AML）專家。請依據使用者提供的 JSON 事實，"
    "以繁體中文撰寫一份「疑似人頭帳戶／洗錢交易分析報告（草稿）」。規定：\n"
    "1. 只能使用 JSON 中的數字與事實，不可杜撰任何未提供的資訊（例如客戶姓名、職業、IP）。\n"
    "2. 章節依序為：一、基本資訊（表格）二、交易概況 三、命中之紅旗指標 四、AI 模型判斷依據 "
    "五、主要資金往來對象 六、研判與稽核建議 七、聲明。\n"
    "3. 第六節請綜合紅旗指標與 AI 依據，說明此帳戶最可能扮演的角色（如收款帳戶、過水帳戶、循環帳戶），"
    "並提出具體可執行的查核步驟。\n"
    "4. 第七節必須包含這段文字：" + DISCLAIMER + "\n"
    "5. 使用 Markdown 格式，語氣專業、精簡。"
)


def _llm_provider(cfg: dict) -> str | None:
    p = cfg.get("provider", "auto")
    if p == "template":
        return None
    if p in ("anthropic", "auto") and os.getenv("ANTHROPIC_API_KEY"):
        return "anthropic"
    if p in ("openai", "auto") and os.getenv("OPENAI_API_KEY"):
        return "openai"
    return None


def _call_anthropic(facts: dict, cfg: dict) -> str:
    model = os.getenv("FLOWAUDIT_ANTHROPIC_MODEL", cfg.get("anthropic_model"))
    r = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={
            "x-api-key": os.environ["ANTHROPIC_API_KEY"],
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json={
            "model": model,
            "max_tokens": 3000,
            "system": LLM_SYSTEM,
            "messages": [{"role": "user", "content": json.dumps(facts, ensure_ascii=False, default=str)}],
        },
        timeout=cfg.get("timeout_sec", 60),
    )
    r.raise_for_status()
    return "".join(b.get("text", "") for b in r.json()["content"])


def _call_openai(facts: dict, cfg: dict) -> str:
    model = os.getenv("FLOWAUDIT_OPENAI_MODEL", cfg.get("openai_model"))
    r = requests.post(
        "https://api.openai.com/v1/chat/completions",
        headers={"Authorization": f"Bearer {os.environ['OPENAI_API_KEY']}"},
        json={
            "model": model,
            "messages": [
                {"role": "system", "content": LLM_SYSTEM},
                {"role": "user", "content": json.dumps(facts, ensure_ascii=False, default=str)},
            ],
        },
        timeout=cfg.get("timeout_sec", 60),
    )
    r.raise_for_status()
    return r.json()["choices"][0]["message"]["content"]


def generate_report(facts: dict, llm_cfg: dict, force_template: bool = False) -> tuple[str, str]:
    """回傳 (報告 Markdown, 實際使用的產生方式)"""
    provider = None if force_template else _llm_provider(llm_cfg)
    if provider:
        try:
            text = _call_anthropic(facts, llm_cfg) if provider == "anthropic" else _call_openai(facts, llm_cfg)
            if DISCLAIMER not in text:
                text += "\n\n## 聲明\n\n" + DISCLAIMER
            return text, provider
        except Exception as e:  # 網路或金鑰問題 → 退回模板
            return render_template(facts) + f"\n\n> 註：LLM 產生失敗（{type(e).__name__}），已改用模板產生。", "template"
    return render_template(facts), "template"


# ---------------------------------------------------------------------------
# 4. 匯出：Word 報告、Excel 工作底稿
# ---------------------------------------------------------------------------
def markdown_to_docx(md: str) -> bytes:
    """把報告 Markdown 轉成 Word（支援標題、表格、清單、粗體）。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Microsoft JhengHei"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")

    def add_runs(par, text):
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for p in parts:
            if p.startswith("**") and p.endswith("**"):
                par.add_run(p[2:-2]).bold = True
            elif p:
                par.add_run(p)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r_i, r in enumerate(rows):
                    for c_i, c in enumerate(r[: len(rows[0])]):
                        cell = t.cell(r_i, c_i)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], c)
                        if r_i == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", line):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\d+\.\s+", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith(">"):
            add_runs(doc.add_paragraph(), line.lstrip("> "))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_workpaper(results: pd.DataFrame, tx: pd.DataFrame, metrics: dict | None, summary: dict,
                    top_k: int = 200, fallback_min: int | None = None, rule_quota: float = 0.0) -> bytes:
    """稽核工作底稿（Excel）：摘要、高風險帳戶清單、規則名單、規則保底名單、規則命中明細、交易明細、模型成效。

    rule_quota > 0 時，top_k 個覆核名額拆成 AI 名單與規則名單（新手法保險，見 model.dual_list）。
    """
    from .model import cycle_fallback, review_lists

    ai_ids, rule_ids = review_lists(results, top_k, rule_quota)
    top = results.loc[ai_ids]
    rl = results.loc[rule_ids]
    fb = (cycle_fallback(results, len(top), fallback_min)
          if fallback_min and {"R4_cycles", "risk_rank"} <= set(results.columns) else results.iloc[:0])
    fb = fb[~fb.index.isin(rl.index)]  # 已在規則名單中的不重複列出
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        info = [
            ("查核程序", "以 FlowAudit 對期間內全部交易進行全查（非抽樣），依 AI 風險分數排序交由人工覆核"),
            ("資料集", summary.get("name")),
            ("帳戶數", summary.get("n_accounts")),
            ("交易筆數", summary.get("n_transactions")),
            ("資料期間", f"{summary.get('date_start')} ~ {summary.get('date_end')}"),
            ("本次覆核帳戶數", f"{len(top) + len(rl)}（AI 名單 {len(top)}、規則名單 {len(rl)}）" if len(rl) else len(top)),
            ("規則保底名單帳戶數", f"{len(fb)}（AI 名單外、參與資金循環 ≥ {fallback_min} 次）" if fallback_min else "未產生"),
            ("底稿產生時間", f"{datetime.now():%Y-%m-%d %H:%M}"),
            ("覆核人員", ""),
            ("覆核日期", ""),
        ]
        if metrics and "model" in metrics:
            info += [
                ("模型 PR-AUC（交叉驗證）", round(metrics["model"]["pr_auc"], 4)),
                ("模型 ROC-AUC（交叉驗證）", round(metrics["model"]["roc_auc"], 4)),
            ]
        pd.DataFrame(info, columns=["項目", "內容"]).to_excel(xw, sheet_name="摘要", index=False)

        cols = {"risk_rank": "排名", "risk_score": "AI風險分數", "risk_level": "風險等級", "rule_hits": "命中規則數",
                "rule_list": "命中規則", "top_reasons": "AI主要判斷依據", "n_in": "匯入筆數", "amt_in": "匯入金額",
                "n_out": "匯出筆數", "amt_out": "匯出金額"}
        lst = top.reset_index()[["account_id"] + [c for c in cols if c in top.columns]].rename(columns={"account_id": "帳戶", **cols})
        lst["覆核結論"] = ""
        lst["備註"] = ""
        lst.to_excel(xw, sheet_name="高風險帳戶清單", index=False)

        if len(rl):
            rcols = {"rule_hits": "命中規則數", "rule_list": "命中規則", "risk_rank": "AI排名", "risk_score": "AI風險分數",
                     "n_in": "匯入筆數", "amt_in": "匯入金額", "n_out": "匯出筆數", "amt_out": "匯出金額"}
            rr = rl.reset_index()[["account_id"] + [c for c in rcols if c in rl.columns]].rename(columns={"account_id": "帳戶", **rcols})
            rr["覆核結論"] = ""
            rr["備註"] = ""
            rr.to_excel(xw, sheet_name="規則名單(新手法保險)", index=False)

        if fallback_min:
            fcols = {"risk_rank": "AI排名", "risk_score": "AI風險分數", "R4_cycles": "參與資金循環次數",
                     "R4_evidence": "循環證據", "rule_list": "命中規則", "amt_in": "匯入金額", "amt_out": "匯出金額"}
            fl = fb.reset_index()[["account_id"] + [c for c in fcols if c in fb.columns]].rename(columns={"account_id": "帳戶", **fcols})
            fl["覆核結論"] = ""
            fl["備註"] = ""
            fl.to_excel(xw, sheet_name="規則保底名單(資金循環)", index=False)

        rows = []
        for a, r in pd.concat([top, rl]).iterrows():
            for code, name in RULES.items():
                if r.get(f"{code}_hit", False):
                    rows.append({"帳戶": a, "規則": f"{code} {name}", "對應態樣／法規": RULE_REFS[code],
                                 "證據": r.get(f"{code}_evidence", "")})
        pd.DataFrame(rows, columns=["帳戶", "規則", "對應態樣／法規", "證據"]).to_excel(xw, sheet_name="規則命中明細", index=False)

        ids = set(top.head(50).index)
        t = tx[tx["src"].isin(ids) | tx["dst"].isin(ids)][["tx_id", "date", "src", "dst", "amount"]].copy()
        t["date"] = t["date"].dt.strftime("%Y-%m-%d")
        t.columns = ["交易序號", "日期", "匯出帳戶", "匯入帳戶", "金額"]
        t.to_excel(xw, sheet_name="交易明細(前50名)", index=False)

        if metrics and "model" in metrics:
            m = []
            for who, key in (("AI 模型", "model"), ("傳統規則", "rules"), ("隨機抽樣", "random_sampling")):
                for r in metrics[key]["at_k"]:
                    m.append({"方法": who, "覆核帳戶數": r["k"], "檢出數": round(r["hits"], 1),
                              "命中率(Precision)": round(r["precision"], 4), "涵蓋率(Recall)": round(r["recall"], 4)})
            pd.DataFrame(m).to_excel(xw, sheet_name="模型成效", index=False)

        for ws in xw.book.worksheets:
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width * 1.6), 80)
    return buf.getvalue()


def top_reasons_text(shap_df: pd.DataFrame, n: int = 3) -> pd.Series:
    s = shap_df.drop(columns=["_bias"], errors="ignore")
    vals = s.to_numpy()
    cols = np.array([FEATURE_NAMES_ZH.get(c, c) for c in s.columns])
    order = np.argsort(-vals, axis=1)[:, :n]
    out = []
    for i in range(len(s)):
        names = [cols[j] for j in order[i] if vals[i, j] > 0]
        out.append("、".join(names))
    return pd.Series(out, index=s.index)
