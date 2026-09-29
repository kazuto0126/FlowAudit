"""產生報告書（Word）。

章節與字級依簡章附件 2「企劃書格式」（壹貳參肆 16 點、一二三 14 點、內文 12 點），封面依去年範例。
文字集中在本檔；表格數字直接讀取結果檔，正文中的關鍵數字在存檔前自動與結果檔核對。
用法：python build_report.py [輸出檔.docx]，接著執行 finalize.ps1 以 Word 更新目錄、移除個人資訊並轉 PDF。
"""
import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
ROOT = HERE.parents[2]
A = HERE / "assets"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "FlowAudit_企劃書.docx"
EA, LAT = "標楷體", "Times New Roman"
BLACK = RGBColor(0, 0, 0)

ev = json.loads((ROOT / "outputs/tw_sim/evaluation.json").read_text(encoding="utf-8"))
_amlw = json.loads((HERE / "amlworld_all.json").read_text(encoding="utf-8"))
ams, AMS_SUM = _amlw["metrics"], _amlw["summary"]
T100 = ev["temporal"]["summary"]["100"]
FAIR = ev["fair"]
_learn = {k: [v["learning"][k] for v in ev["unseen"]["schemes"].values()] for k in ("0", "10")}
UNSEEN = f"{min(_learn['0']):.0%}～{max(_learn['0']):.0%}"      # 沒見過的手法，AI 找回該類的比例
LEARN10 = f"{min(_learn['10']):.0%}～{max(_learn['10']):.0%}"   # 新手法累積 10 個警示帳戶後
_tx = pd.read_csv(ROOT / "data/raw/tw_sim/transactions.csv.gz", usecols=["src", "pattern"])
N_VICTIMS = _tx.loc[_tx["pattern"] == "被害人匯入", "src"].nunique()  # 實際有匯款的被害人
_acc = pd.read_csv(ROOT / "data/raw/tw_sim/accounts.csv", usecols=["label", "mule_layer", "alert_date"])
LAYER = {int(k): (len(g), int(g["alert_date"].notna().sum()))   # 各層人頭帳戶數與期間內被警示數
         for k, g in _acc[_acc["label"] == 1].groupby("mule_layer")}

doc = Document()


# ------------------------------------------------------------------ 樣式
def set_fonts(el_rpr):
    rfonts = el_rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        el_rpr.insert(0, rfonts)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]
    rfonts.set(qn("w:ascii"), LAT)
    rfonts.set(qn("w:hAnsi"), LAT)
    rfonts.set(qn("w:cs"), LAT)
    rfonts.set(qn("w:eastAsia"), EA)


def style(name, size, bold=False, before=0, after=4, line=1.2, indent=None, align=None, keep_next=False):
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.color.rgb = BLACK
    set_fonts(s.element.get_or_add_rPr())
    pf = s.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if indent is not None:
        pf.first_line_indent = indent
    if align is not None:
        pf.alignment = align
    pf.keep_with_next = keep_next
    return s


# 字級依簡章附件 2：壹貳參肆 16 點、一二三 14 點、內文 12 點
style("Normal", 12, indent=Pt(24), align=WD_ALIGN_PARAGRAPH.JUSTIFY)
style("Heading 1", 16, bold=True, before=0, after=10, line=1.0, indent=Pt(0), keep_next=True)
style("Heading 2", 14, bold=True, before=10, after=4, line=1.0, indent=Pt(0), keep_next=True)
style("Heading 3", 12, bold=True, before=6, after=2, line=1.0, indent=Pt(0), keep_next=True)
style("Heading 4", 12, bold=True, before=4, after=2, line=1.0, indent=Pt(0), keep_next=True)
style("List Bullet", 12, indent=None, after=2)
style("Caption", 10.5, before=2, after=8, line=1.0, indent=Pt(0), align=WD_ALIGN_PARAGRAPH.CENTER)

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.3)
sec.left_margin = sec.right_margin = Cm(2.3)


# ------------------------------------------------------------------ 內容工具
def runs(p, text, size=None, bold_all=False):
    """**粗體** 以外的文字照原樣輸出。"""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        if bold_all or i % 2 == 1:  # 不明確設定為「不粗體」，以免蓋掉標題樣式的粗體
            r.bold = True
        if size:
            r.font.size = Pt(size)
    return p


def P(text, indent=True, size=None, align=None, after=None, before=None):
    p = doc.add_paragraph()
    if not indent:
        p.paragraph_format.first_line_indent = Pt(0)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    return runs(p, text, size)


def H(level, text, page_break=False):
    h = doc.add_heading(level=level)
    runs(h, text)
    if page_break:
        h.paragraph_format.page_break_before = True
    return h


def numbered(items, fmt="{i}.　"):
    """編號清單自行編號（內建編號樣式會跨清單連續編號）。"""
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.space_after = Pt(2)
        runs(p, fmt.format(i=i) + t)


def bullets(items):
    for t in items:
        runs(doc.add_paragraph(style="List Bullet"), t)


# 圖表編號依出現順序；正文用 F("key")、T("key") 引用，建立時檢查順序一致
FIGS = ["arch", "workflow", "shot_overview", "shot_list", "shot_investigate", "shot_evidence", "shot_group",
        "equal", "solutions", "external"]
TABS = ["features", "data", "usage", "rules", "featgroups", "lists", "workpaper", "equal", "temporal", "delay",
        "dilemmas", "amlworld", "apps"]
FIG, TAB = [0], [0]


def F(key):
    return f"圖 {FIGS.index(key) + 1}"


def T(key):
    return f"表 {TABS.index(key) + 1}"


def figure(key, path, width_cm, caption):
    assert FIGS[FIG[0]] == key, (key, FIGS[FIG[0]])
    FIG[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = doc.add_paragraph(style="Caption")
    runs(c, f"{F(key)}　{caption}")


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(key, caption, rows, widths_cm, size=10, header=True, bold_first_col=False, align_right=()):
    assert TABS[TAB[0]] == key, (key, TABS[TAB[0]])
    TAB[0] += 1
    c = doc.add_paragraph(style="Caption")
    c.paragraph_format.keep_with_next = True
    c.paragraph_format.space_after = Pt(2)
    runs(c, f"{T(key)}　{caption}")
    t = doc.add_table(rows=len(rows), cols=len(widths_cm))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Cm(widths_cm[j])
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.paragraph_format.keep_with_next = i < len(rows) - 1
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if (j in align_right and i > 0) else WD_ALIGN_PARAGRAPH.LEFT
            runs(p, str(val), size=size, bold_all=(header and i == 0) or (bold_first_col and j == 0))
            if header and i == 0:
                shade(cell, "E9F1FB")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.6
    return t


def callout(title, items, size=10.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(16.4)
    shade(cell, "F4F4F2")
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    runs(p, title, size=size + 0.5, bold_all=True)
    for i, it in enumerate(items, 1):
        q = cell.add_paragraph()
        q.paragraph_format.first_line_indent = Pt(0)
        q.paragraph_format.left_indent = Pt(14)
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing = 1.15
        runs(q, f"{i}. {it}", size=size)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.line_spacing = 0.6


def field(paragraph, instr, placeholder=""):
    r = paragraph.add_run()
    for kind, txt in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = txt
            r._r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def wan(x):
    """元 → 萬元（四捨五入、千分位）。"""
    return f"{x / 1e4:,.0f}"


# 正文關鍵數字與結果檔核對（存檔前檢查）
CHECKS = []


def check(label, shown, actual):
    CHECKS.append((label, shown, actual))


# ================================================================== 封面（依去年範例）
def blank(n, size=12):
    for _ in range(n):
        P("", indent=False, size=size, after=0)


blank(3)
P("2026 法遵科技與電腦稽核專題競賽", indent=False, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
P("企劃書", indent=False, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
cover = doc.add_table(rows=1, cols=1)
cover.style = "Table Grid"
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
cc = cover.cell(0, 0)
cc.width = Cm(15.0)
p = cc.paragraphs[0]
p.paragraph_format.first_line_indent = Pt(0)
p.paragraph_format.space_before = Pt(10)
runs(p, "作品名稱：", size=16)
for txt, size, align in (("", 16, None), ("FlowAudit 金流偵探：以資金流向圖與可解釋 AI", 18, WD_ALIGN_PARAGRAPH.CENTER),
                         ("偵測人頭帳戶的持續稽核系統", 18, WD_ALIGN_PARAGRAPH.CENTER), ("", 12, None)):
    q = cc.add_paragraph()
    q.paragraph_format.first_line_indent = Pt(0)
    q.paragraph_format.line_spacing = 1.3
    if align is not None:
        q.alignment = align
    runs(q, txt, size=size)
blank(21)
P("中華民國 115 年 12 月 04 日", indent=False, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

# ================================================================== 目錄
p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Pt(0)
p.add_run().add_break(WD_BREAK.PAGE)
P("**目錄**", indent=False, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
toc = doc.add_paragraph()
toc.paragraph_format.first_line_indent = Pt(0)
field(toc, 'TOC \\o "1-3" \\h \\z \\u', "（開啟檔案後在此按右鍵 → 更新功能變數，即可產生目錄）")

# 正文另起一節，頁碼從 1 開始
body = doc.add_section(WD_SECTION.NEW_PAGE)
pg = OxmlElement("w:pgNumType")
pg.set(qn("w:start"), "1")
body._sectPr.append(pg)
body.footer.is_linked_to_previous = False
fp = body.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Pt(0)
field(fp, "PAGE", "1")
doc.sections[0].footer.is_linked_to_previous = False

# ================================================================== 壹、前言
H(1, "壹、前言")
H(2, "一、作品主題")
H(3, "（一）作品簡介")
P("FlowAudit（金流偵探）是一套偵測人頭帳戶的持續稽核系統。它對期間內的**每一筆交易、每一個帳戶全查**："
  "先以 8 條對應法規態樣的紅旗規則檢核，再把帳戶之間的資金往來建成「資金流向圖」，萃取 57 項特徵交給 AI 評分，"
  "並逐一說明每個帳戶為什麼可疑；最後自動產生覆核名單、可疑交易分析報告與 Excel 稽核工作底稿，"
  "讓稽核人員把有限的人力用在最可能有問題的帳戶上。")

H(3, "（二）背景")
P("詐騙金流離不開人頭帳戶。2025 年全台詐騙財損約 893 億元、受理約 16.2 萬件 [4]；"
  "截至 2026 年 1 月，全台仍有約 14.85 萬個帳戶被列為警示帳戶 [5]。詐騙集團收到被害人款項後，"
  "會透過第一層、第二層人頭帳戶在數小時內層層轉出，再由車手提領或轉入虛擬貨幣交易所 [6][7]。"
  "人頭帳戶是詐騙金流中不可缺少的一環，也是金融機構最有機會攔截的環節。")

H(3, "（三）動機")
any_ = FAIR["any_rule"]
P("《存款帳戶及其疑似不法或顯屬異常交易管理辦法》要求銀行辨識與處理疑似不法或顯屬異常的帳戶 [2]，"
  "法務部調查局也訂有銀行業「疑似洗錢或資恐交易態樣」供申報參考 [1]。依態樣設定門檻的監控規則有法規依據、"
  "容易解釋，但仍有三個限制，是本作品想改善的地方：")
numbered([f"**固定門檻誤報多**：行為相似的正常帳戶會一併被列為可疑。本研究資料中，若把命中任一規則的帳戶都列為可疑，"
          f"共 {any_['n']:,} 個帳戶，其中 {any_['normal']:,} 個是正常帳戶；團購主、發薪公司、房東這類行為像人頭帳戶的"
          "正常帳戶幾乎全部被列入。",
          "**多在事後才發現**：警示帳戶是在警方等機關通報後才列入 [2]，此時詐騙款項多已層層轉出 [6]。",
          "**抽樣看不到關係**：內部稽核以抽樣測試監控機制時，人頭帳戶的特徵藏在帳戶之間的關係與資金移動的速度裡，"
          "只看抽出的交易不易發現。"])

H(3, "（四）目的")
P("本作品希望做到**全查不抽樣、比警方通報更早發現、說得出為什麼可疑、結果能直接成為稽核工作底稿**；"
  "並誠實驗證這套方法在什麼情況下有效、在什麼情況下會失效，清楚區分「已在仿真資料上證實」與「仍屬推估」的結果。")

H(3, "（五）研究主題目前的沿革")
P("人頭帳戶與洗錢偵測的做法，大致可分為三個階段：")
numbered(["**規則式監控**：依法規與態樣設定門檻，例如調查局的疑似洗錢態樣 [1] 與管理辦法的異常交易類型 [2]。"
          "優點是有法規依據、容易解釋；缺點是門檻固定，行為相似的正常帳戶容易被誤報，也不易涵蓋新手法。",
          "**機器學習評分**：以已知案例訓練模型、依風險高低排序，讓有限的調查人力先看最可疑的帳戶；但模型需要有答案的"
          "訓練資料，也必須說得出判斷依據，才能符合金融業運用 AI 重視的透明性與可解釋性 [8]。TreeSHAP 等方法已能逐一"
          "說明樹狀模型對每一筆資料的判斷依據 [12]。",
          "**資金網路分析**：把交易視為帳戶之間的網路，找出集團、資金循環等單看一個帳戶看不出的結構 [13]；"
          "IBM 等研究團隊並公開合成的反洗錢交易資料，供網路分析與圖神經網路等方法比較 [9][10]。"])
P("本作品把三者整合成稽核人員每月可以使用的流程：規則保留法規依據與證據，AI 以規則背後的連續數值與其他行為特徵排序，"
  "資金流向圖用於集團調查與循環舉證；並以「只用已警示帳戶訓練」「集團層級交叉驗證」「逐月模擬」等貼近實務的方式評估。")

H(2, "二、作品特色")
H(3, "（一）創意概念")
numbered(["**看「關係」與「速度」，不只看單筆金額**：把交易建成資金流向圖，偵測快進快出、集中匯入、分散匯出、"
          "時間先後一致的資金循環等人頭帳戶典型行為。",
          "**規則與 AI 分工，而不是互相取代**：規則有法規依據、容易解釋；AI 能從已警示的帳戶學習，但只學得到看過的手法。"
          "本作品把覆核名額分成「AI 名單、規則名單、規則保底名單」三份，讓兩者互補。",
          "**持續稽核**：每月全查一次，稽核人員的覆核結論回饋給模型，下個月判斷更準；新的詐騙手法只要累積少數警示帳戶，"
          "模型就能學會。"])
H(3, "（二）特色、解決的問題與提供的便利")
eq = {r["k"]: r for r in FAIR["by_k"]}
n_pos = FAIR["n_mules"]
r80 = next(r for r in FAIR["equal_recall"] if r["recall"] == 0.8)
P(f"{T('features')} 整理本作品的特色。表中數字皆來自本研究的台灣情境仿真資料（資料來源見「貳、一、（一）」），"
  "「規則」指本研究設定的比較基準（見「貳、一、（四）」）。")
table("features", "作品特色、解決的問題與提供的便利", [
    ["特色", "解決的問題", "提供的便利"],
    ["全查＋依風險排序", "調查人力有限，只能看一部分帳戶",
     f"每月只覆核 100 個帳戶，6 個月找出 {T100['ai']['unique_mules']} 個人頭帳戶（規則 {T100['rules']['unique_mules']} 個）"],
    ["相同人力下誤報少", "門檻固定，正常帳戶被反覆調查",
     f"覆核同樣 {n_pos} 個帳戶，名單中的正常帳戶 {eq[n_pos]['ai']['normal']:.0f} 個（規則 {eq[n_pos]['rules']['normal']:.0f} 個）"],
    ["提早發現", "帳戶多在警方通報後才被警示",
     f"逐月模擬中 {T100['ai']['caught_before_alert']} 個比警方通報平均早 {T100['ai']['lead_days_mean']:.0f} 天"],
    ["有所本、說得出為什麼", "AI 判斷難以向主管與受查者說明", "每條規則對應態樣或法規並附證據；逐帳戶列出 AI 判斷依據"],
    ["冷啟動與持續學習", "導入初期警示資料少、新手法不斷出現",
     f"冷啟動 6 個月找到 {T100['cold']['unique_mules']} 個；新手法累積 10 個警示帳戶後，找回率回升到 {LEARN10}"],
    ["自動產出工作底稿", "調查結果整理費時、不易追溯", "一鍵匯出可疑交易分析報告與 Excel 工作底稿，可完全離線運作"],
], [3.6, 5.2, 7.6], bold_first_col=True)

# ================================================================== 貳、成果
H(1, "貳、成果", page_break=True)
H(2, "一、摘要")
P(f"本作品完成一套可實際操作的持續稽核系統，包括紅旗規則引擎、特徵工程與資金流向圖、AI 風險模型與可解釋分析、"
  f"三份覆核名單、7 頁的稽核儀表板，以及可疑交易分析報告與 Excel 工作底稿的自動產出。在台灣情境仿真資料上，"
  f"「每月覆核 100 個帳戶」的逐月模擬中，AI 6 個月找出 {T100['ai']['unique_mules']} 個人頭帳戶，比較基準的規則為 "
  f"{T100['rules']['unique_mules']} 個；覆核同樣多的帳戶時，AI 名單中的正常帳戶也明顯較少。本研究同時測試 AI 會失效的情境"
  "（導入初期、新手法、循環交易）並提出因應，另在 IBM 公開的合成資料 AMLworld 上以其標籤重新訓練，驗證方法仍然有效。"
  "以下依序說明資料來源、系統架構、系統展示與成效驗證，並標明哪些結果已在仿真資料上證實、哪些仍屬推估。")

# ------------------------------------------------------------------ （一）資料
H(3, "（一）資料來源與取得方式")
P("真實銀行交易受個人資料保護法與銀行保密規定保護，學生無法取得。因此本研究的訓練與主要評估資料，"
  f"是以自行撰寫的「台灣情境交易仿真器」產生的合成資料（{T('data')}），不含任何真實個人資料。取得方式分為四個步驟：")
numbered(["**蒐集依據**：依調查局疑似洗錢態樣 [1]、金管會管理辦法 [2]、165 打詐儀錶板統計 [4]、新聞報導 [6][7] "
          "與銀行公開的轉帳限額 [14]，決定詐騙金流的樣態與參數。",
          "**程式產生**：以 Python 撰寫的仿真器產生 6 個月的帳戶與交易。仿真器固定隨機種子，任何人執行同一個指令都能"
          "產生逐筆相同的資料。為了避免模型只學到仿真器的破綻，仿真器刻意加入大量**行為像人頭帳戶、但其實完全正常**的帳戶，"
          "例如團購主、標會會頭、個人賣家、同日匯薪給數十名員工的公司、休眠後領保險金的帳戶，以及家人共用手機、"
          "記帳士代管多家公司網銀等情形。",
          "**產生「答案」**：答案不是人工標註，而是模擬銀行的警示流程——被害人依詐騙類型過一段時間才報案（假客服解除分期"
          " 1～3 天、網購詐騙 1～10 天、假投資在最後一次匯款後 3～30 天），收款的第一層人頭帳戶在第一位被害人報案當天被列為"
          "警示帳戶並凍結，集團隨即改用下一個帳戶；第二層帳戶約六成在開始使用後 20～70 天被追查到；循環交易沒有被害人，"
          "只有約兩成會被警示，且多在觀察期之後。結果第一層人頭帳戶 187 個中有 186 個在期間內被警示、第二層 39 個中 27 個、"
          "循環交易 17 個中只有 2 個，合計 215 個被警示、28 個從未被警示。",
          "**訓練取用**：模型的輸入只有交易資料與銀行本來就知道的帳戶屬性（開戶日、觀察期前未往來天數、客戶類型）；"
          "答案只用「期間內是否被警示」，這是銀行實際上唯一知道的答案。因此 28 個從未被警示的人頭帳戶在訓練時被當成"
          "正常帳戶，評估時則算人頭帳戶，用來檢驗模型能否找出「還沒被通報」的帳戶。仿真器知道的真實角色、所屬集團與"
          "詐騙類型只用於評分，不提供給模型。"])
check("人頭帳戶總數", 243, ev["n_true_positive"])
check("已警示人頭帳戶", 215, ev["n_train_positive"])
check("第一層（總數、被警示）", (187, 186), LAYER[1])
check("第二層（總數、被警示）", (39, 27), LAYER[2])
check("循環交易（總數、被警示）", (17, 2), LAYER[3])
table("data", "台灣情境仿真資料", [
    ["項目", "內容"],
    ["期間", "2026 年 1 月 1 日至 6 月 30 日（6 個月）"],
    ["規模", f"{ev['n_accounts']:,} 個帳戶、{len(_tx):,} 筆交易（含現金存提）；被害人 {N_VICTIMS} 名、匯入約 "
           f"{wan(ev['temporal']['summary']['victim_amount_total'])} 萬元"],
    ["人頭帳戶", f"{ev['n_true_positive']} 個（1.4%），來自 20 個詐騙集團；{ev['n_train_positive']} 個在期間內被警示"
             "（第一層 186／187、第二層 27／39、循環交易 2／17），28 個從未被警示"],
    ["詐騙手法", "假投資、網購詐騙、假客服解除分期、循環交易"],
    ["欄位", "交易時間、金額（新台幣）、交易管道、登入裝置、開戶日、觀察期前未往來天數、客戶類型"],
    ["設計依據", "金流「被害人 → 第一層 → 第二層 → 車手提領或虛擬貨幣」依公開報導 [6][7]；ATM 單次 2 萬元、"
             "每日 15 萬元等限額依銀行公開資訊 [14]"],
], [2.6, 13.8], bold_first_col=True)
P(f"本研究用到的所有資料與用途整理於 {T('usage')}。除了 IBM AMLworld 之外都出自同一個仿真器，因此另以 AMLworld "
  "檢查方法在不同來源資料上是否仍然有效；真實銀行資料上的成效則尚待驗證。實務導入時，訓練資料直接換成機構自己的"
  "交易紀錄與警示帳戶名單即可，不需要另外請人標註。")
table("usage", "資料用途一覽", [
    ["資料", "來源與取得方式", "本研究的用途", "是否重新訓練"],
    ["台灣情境仿真資料", "本研究仿真器產生（隨機種子 20260101）", "模型訓練與主要評估：集團層級交叉驗證、逐月模擬", "以此訓練"],
    ["另一份仿真資料", "同一仿真器、另一個隨機種子（20260102）", "仿真穩健性測試；冷啟動的參考資料",
     "穩健性測試不重新訓練；冷啟動以其已警示帳戶訓練參考模型"],
    ["壓力測試資料（6 種情境）", "同一仿真器，調整人頭帳戶的行為參數（隨機種子 20260103）",
     "仿真穩健性測試；「以上全部」情境另作為分布不同的冷啟動參考資料", "重新訓練與直接套用兩種都測"],
    ["IBM AMLworld（HI-Small）", "Kaggle 公開下載的 IBM 合成資料 [9][15]；記錄檔案雜湊值（SHA-256）以驗證檔案一致",
     "檢查方法在不同來源的資料上是否仍然有效", "是：以該資料本身的標籤重新訓練，5 折交叉驗證"],
    ["真實銀行資料", "尚未取得", "導入時以機構自己的交易與警示紀錄重新訓練與驗證", "—"],
], [3.0, 5.0, 4.6, 3.8], size=9.5, bold_first_col=True)

# ------------------------------------------------------------------ （二）架構
H(3, "（二）系統架構")
P(f"系統架構如{F('arch')}。交易資料依序經過五個模組，每個模組的結果都能單獨檢視與追溯；稽核人員的覆核結論會回饋給模型，"
  "形成每月的持續稽核循環。")
figure("arch", A / "fig_architecture.png", 14.5, "系統架構")

H(4, "1. 紅旗規則引擎")
P(f"規則引擎對每一個帳戶檢核 8 條紅旗規則（{T('rules')}），每條都對應調查局疑似洗錢態樣代碼 [1] 或管理辦法條文 [2]，"
  "命中時附上證據文字，例如「共 9 筆流入在 24 小時內被轉出 80% 以上」。門檻依客群（個人、商家、公司）分別計算，"
  "並集中在設定檔中，稽核人員可依機構風險胃納調整、不需修改程式；資料缺少時間、現金、裝置或開戶欄位時，"
  "對應規則會自動改用替代判斷或不觸發。")
table("rules", "紅旗規則與法規依據", [
    ["代碼", "規則", "對應態樣／法規", "判斷方式（預設門檻）"],
    ["R1", "快進快出", "態樣 A15、A16、A17", "流入後 24 小時內轉出或提領 ≥ 80%，至少 2 次"],
    ["R2", "集中匯入", "態樣 A17", "7 日內不同匯入來源數 ≥ 該客群第 95 百分位"],
    ["R3", "分散匯出", "態樣 A18", "7 日內不同匯出對象數 ≥ 該客群第 95 百分位"],
    ["R4", "資金循環", "態樣 A18", "依時間先後、金額相近地經 2～5 個帳戶回流，30 日內完成"],
    ["R5", "休眠／新戶突然活躍", "態樣 A15、A16；管理辦法第 4 條", "開戶 ≤ 90 天或未往來 ≥ 180 天，開始使用 14 日內流入 ≥ 5 萬元且轉出 ≥ 70%"],
    ["R6", "密集現金提領", "態樣 A11、A12", "24 小時內 ATM／臨櫃提領 ≥ 3 次且 ≥ 5 萬元"],
    ["R7", "多帳戶共用裝置", "態樣 A1G", "同一登入裝置有 ≥ 3 個帳戶轉入同一受款帳號"],
    ["R8", "小額測試交易", "管理辦法第 4 條", "72 小時內 ≥ 3 筆 30 元以下的轉出入"],
], [1.3, 3.2, 4.4, 7.5])

H(4, "2. 特徵工程與資金流向圖")
P(f"規則只回答「有沒有超過門檻」，會丟掉許多資訊。本作品保留規則背後的連續數值，並從資金流向圖計算每個帳戶"
  f"在網路中的位置，共 57 項特徵、分為 6 類（{T('featgroups')}）。特徵只使用交易資料與帳戶屬性，"
  "不使用任何「答案」資訊，避免評估結果虛高。")
table("featgroups", "特徵類別", [
    ["類別", "項數", "例子"],
    ["交易行為", "17", "匯入匯出筆數與金額、不重複往來對象數、首末交易間隔天數、整千元交易比例"],
    ["資金流速", "5", "流入後至下一筆轉出的平均小時數、流入後 24 小時內即轉出的比例、快進快出金額占比"],
    ["時間與管道", "9", "深夜交易比例、現金提領占轉出比例、24 小時內最多提領次數"],
    ["資金網路圖", "14", "7 日內最多不同匯入來源、參與資金循環次數、資金網路重要性（PageRank）、往來對象命中規則比例"],
    ["帳戶屬性與裝置", "11", "首次交易時的開戶天數、未往來天數、登入裝置共用帳戶數、極小額交易筆數"],
    ["規則命中數", "1", "命中紅旗規則的條數"],
], [3.1, 1.3, 12.0])

H(4, "3. AI 風險模型與可解釋性")
P("風險模型採用 XGBoost 梯度提升樹 [11]，訓練時**只用「已被通報警示」的帳戶**當正樣本（見（一））。每個帳戶的風險分數"
  "都來自沒看過它的模型（交叉驗證的樣本外分數），分數 0.8 以上為高風險、0.5 以上為中風險。可解釋性方面，本作品以 "
  "TreeSHAP [12] 精確拆解每個帳戶的分數：哪些特徵把風險往上推、推了多少，例如「首末交易間隔只有 4 天」"
  "「流入後 24 小時內即轉出的比例 100%」。搭配規則的證據文字，稽核人員能說清楚「為什麼這個帳戶可疑」。此外，系統在"
  "中、高風險帳戶的資金往來子圖上以 Louvain 社群偵測 [13] 切出疑似集團，協助以「集團」為單位擴大查核。")

H(4, "4. 覆核名單與冷啟動")
P(f"AI 只學得到見過、而且有人被通報過的手法。因此本作品不把覆核名額全部交給 AI，而是分成三份名單（{T('lists')}）；"
  "比例與門檻都可在設定檔調整。")
table("lists", "三份覆核名單", [
    ["名單", "內容", "目的"],
    ["AI 名單", "覆核名額的 90%，依 AI 風險分數由高到低", "找出與已知案例相似的人頭帳戶"],
    ["規則名單", "覆核名額的 10%，規則分數最高、但不在 AI 名單的帳戶", "新詐騙手法出現時的保險"],
    ["規則保底名單", "名單以外、參與時間一致資金循環 ≥ 5 次的帳戶，另案專案查核", "補上很少被通報、AI 學不到的循環交易"],
], [2.6, 7.6, 6.2], bold_first_col=True)
P("導入初期，機構自己的警示帳戶很少，模型學不起來。**冷啟動**的做法是：自家已知警示帳戶不到 50 個時，先使用以外部參考資料"
  "訓練的模型，累積足夠的警示帳戶後再改用自家模型。實務上參考資料可來自同業或主管機關分享、已完成調查的資料；"
  "**本研究則以同一仿真器、另一個隨機種子產生的仿真資料代表參考資料**（另以分布不同的壓力測試資料做較保守的比較），"
  "參考資料同樣只用已警示帳戶當正樣本。")

# ------------------------------------------------------------------ （三）展示
H(3, "（三）系統展示與使用情境")
P("FlowAudit 的主要使用者是銀行的法遵暨洗錢防制單位與內部稽核單位，對應內部控制的「三道防線」：")
bullets([
    "**第一道防線（營業單位）**：開戶與臨櫃審核時，可參考帳戶的風險等級與命中規則。",
    "**第二道防線（法遵、洗錢防制）**：每月依覆核名單調查帳戶，決定是否列為警示帳戶或申報疑似洗錢交易。",
    "**第三道防線（內部稽核）**：以 FlowAudit 獨立全查，檢驗既有監控機制是否漏掉可疑帳戶，作為控制有效性測試的佐證。",
])
P(f"此外，會計師事務所在執行查核或洗錢防制諮詢時，也可用「上傳新資料全查」功能對客戶提供的交易資料進行全查。"
  f"每月的作業流程如{F('workflow')}。")
figure("workflow", A / "fig_workflow.png", 15.5, "每月持續稽核作業流程")

H(4, "1. 稽核儀表板")
P(f"儀表板共 7 頁：總覽、高風險帳戶清單、帳戶調查、疑似集團分析、模型評估、上傳新資料全查、系統架構與設定。"
  f"{F('shot_overview')} 至{F('shot_group')} 為主要畫面，為便於閱讀只擷取關鍵區塊。畫面中的「仿真真實角色」等欄位是"
  "仿真資料才有的驗證資訊，實務上不會出現。")
figure("shot_overview", A / "shot_overview.png", 15.5, "總覽：全查規模與逐月持續稽核模擬的成果摘要")
figure("shot_list", A / "shot_list.png", 14.5,
       "高風險帳戶清單（前 8 名）與另外兩份覆核名單；清單另有「AI 主要判斷依據」欄，可一鍵匯出工作底稿")
figure("shot_investigate", A / "shot_investigate.png", 14,
       "帳戶調查：左為 AI 判斷依據（紅色推高風險、藍色降低風險），右為該帳戶的資金往來網路")
figure("shot_evidence", A / "shot_evidence.png", 15.5, "帳戶調查：命中紅旗規則的態樣代碼與證據文字")
figure("shot_group", A / "shot_group.png", 12.5, "疑似集團分析：以社群偵測切出的集團資金網路與成員名單")

H(4, "2. 可疑交易分析報告與稽核工作底稿")
P("在帳戶調查頁按下「產生報告」，系統會產生可疑交易分析報告草稿，內容包括基本資訊、交易概況、命中的紅旗指標（附態樣與證據）、"
  "AI 判斷依據、主要資金往來對象、稽核建議與聲明，可匯出 Word。預設以模板產生、完全離線；若設定語言模型金鑰，"
  f"也可改由語言模型撰寫，但只提供結構化事實並要求不得杜撰，失敗時自動退回模板。稽核工作底稿則匯出為 Excel（{T('workpaper')}），"
  "保留覆核結論與備註欄位，覆核人員填寫後即可歸檔。")
table("workpaper", "稽核工作底稿的工作表", [
    ["工作表", "內容"],
    ["摘要", "查核程序、資料期間、覆核帳戶數、模型成效，以及覆核人員與覆核日期欄位"],
    ["高風險帳戶清單", "AI 名單：排名、風險分數、命中規則、AI 主要判斷依據、覆核結論與備註"],
    ["規則名單（新手法保險）", "規則分數最高、但不在 AI 名單的帳戶"],
    ["規則保底名單（資金循環）", "參與資金循環的次數與循環證據"],
    ["規則命中明細", "每一條命中規則的對應態樣／法規與證據文字"],
    ["交易明細、模型成效", "前 50 名帳戶的相關交易；各覆核名額下 AI、規則與抽樣的命中率"],
], [4.6, 11.8], bold_first_col=True)

# ------------------------------------------------------------------ （四）驗證
H(3, "（四）成效驗證")
H(4, "1. 評估設計與比較基準")
P("為了避免高估成效，本研究採用下列設計：（1）**只用已警示帳戶訓練**，以全部人頭帳戶評估；（2）**集團層級交叉驗證**，"
  "同一詐騙集團的帳戶不會同時出現在訓練與測試，避免模型只是記住某個集團；（3）**逐月持續稽核模擬**，每月 1 日只用當時"
  "已有的交易與已知的警示帳戶；（4）**大量困難的正常帳戶**（見（一））。")
P("**比較基準**為「現行規則計分」：以同一套 8 條紅旗規則，依命中條數排序、同分時再依快進快出金額占比排序，代表不使用"
  "機器學習、只依規則排定調查順序的做法。它是本研究設定的基準，不代表所有機構的實際做法。主要指標 PR-AUC 衡量"
  "「依風險由高到低覆核時，能不能把人頭帳戶排在前面」，介於 0 到 1、越高越好，隨機抽樣約等於人頭帳戶的比例"
  "（本資料約 0.014）。PR-AUC 是整條排序的綜合指標，不等於找回比例，因此以下同時列出「覆核多少帳戶、找到多少人頭帳戶」。")

H(4, "2. 相同覆核量下的比較")
P(f"公平比較的前提是**覆核同樣多的帳戶**：AI 與規則都依各自的分數由高到低覆核，比較找到的人頭帳戶與名單中的正常帳戶"
  f"（{T('equal')}、{F('equal')}）。AI 分數為集團層級交叉驗證的樣本外分數；規則分數同分時，覆核順序視為隨機並取期望值。")
rows = [["覆核帳戶數", "AI：找到人頭帳戶", "AI：名單中的正常帳戶", "規則：找到人頭帳戶", "規則：名單中的正常帳戶", "隨機抽樣（期望值）"]]
for r in FAIR["by_k"]:
    k = r["k"]
    lab = f"{k:,}（＝人頭帳戶總數）" if k == n_pos else f"{k:,}"
    rows.append([lab, f"{r['ai']['mules']:.0f}", f"{r['ai']['normal']:,.0f}", f"{r['rules']['mules']:.0f}",
                 f"{r['rules']['normal']:,.0f}", f"{r['random_mules']:.1f}"])
rows.append([f"{any_['n']:,}（＝命中任一規則的帳戶數）", f"{any_['ai_same_k']['mules']:.0f}", f"{any_['ai_same_k']['normal']:,.0f}",
             f"{any_['mules']}", f"{any_['normal']:,}", f"{any_['n'] * n_pos / FAIR['n_accounts']:.1f}"])
table("equal", "相同覆核量下，AI 與比較基準（規則）找到的人頭帳戶與名單中的正常帳戶", rows,
      [4.4, 2.3, 2.5, 2.3, 2.5, 2.4], size=9.5, align_right=(1, 2, 3, 4, 5))
figure("equal", A / "fig_equal_volume.png", 15.5, "相同覆核量下累計找到的人頭帳戶（台灣情境仿真資料）")
r90 = next(r for r in FAIR["equal_recall"] if r["recall"] == 0.9)
nr = FAIR["normal_by_role_at_n_pos"]
P(f"覆核同樣 {n_pos} 個帳戶（等於人頭帳戶總數）時，AI 找到 {eq[n_pos]['ai']['mules']:.0f} 個人頭帳戶、名單中有 "
  f"{eq[n_pos]['ai']['normal']:.0f} 個正常帳戶；規則找到 {eq[n_pos]['rules']['mules']:.0f} 個、正常帳戶 "
  f"{eq[n_pos]['rules']['normal']:.0f} 個。換成「找到同樣比例」來比：要找回 80% 的人頭帳戶（{r80['mules']} 個），"
  f"AI 需覆核 {r80['ai']['reviewed']:,.0f} 個帳戶，規則需覆核 {r80['rules']['reviewed']:,.0f} 個（其中 "
  f"{r80['rules']['normal']:,.0f} 個是正常帳戶）；要找回 90%，AI 需覆核 {r90['ai']['reviewed']:,.0f} 個，規則約需 "
  f"{round(r90['rules']['reviewed'], -2):,.0f} 個。若把命中任一規則的 {any_['n']:,} 個帳戶全部列為可疑，雖涵蓋 "
  f"{any_['mules']} 個人頭帳戶，卻要同時調查 {any_['normal']:,} 個正常帳戶。")
P(f"名單中的正常帳戶是哪些？覆核前 {n_pos} 名時，AI 誤列的 {eq[n_pos]['ai']['normal']:.0f} 個主要是學生（{nr['ai'].get('學生', 0)} 個）、"
  f"代收款項帳戶（{nr['ai'].get('代收帳戶', 0)} 個）與第二帳戶（{nr['ai'].get('第二帳戶', 0)} 個）；規則誤列的 "
  f"{eq[n_pos]['rules']['normal']:.0f} 個主要是個人賣家（{nr['rules'].get('個人賣家', 0)} 個）與上班族（{nr['rules'].get('上班族', 0)} 個）。"
  "若改以命中任一規則為準，團購主與代收款項帳戶 100%、發薪公司 99%、房東 96% 都會被列入，這些正是「行為像人頭帳戶、"
  "其實正常」的客群。")

H(4, "3. 逐月持續稽核模擬與潛在攔阻金額")
P(f"每月 1 日執行一次全查，對尚未被警示的帳戶排序並覆核風險最高的 100 個，6 個月共覆核 600 個帳戶，結果如{T('temporal')}。")
table("temporal", "逐月持續稽核模擬（每月覆核 100 個帳戶，6 個月合計）", [
    ["指標", "AI 模型", "比較基準：規則", "隨機抽樣"],
    ["找出的人頭帳戶（不重複）", f"**{T100['ai']['unique_mules']} 個**", f"{T100['rules']['unique_mules']} 個",
     f"約 {T100['random']['unique_mules']:.0f} 個"],
    ["比警方通報更早發現", f"**{T100['ai']['caught_before_alert']} 個，平均早 {T100['ai']['lead_days_mean']:.0f} 天**",
     f"{T100['rules']['caught_before_alert']} 個，平均早 {T100['rules']['lead_days_mean']:.0f} 天", "—"],
    ["期間內從未被通報、由系統找出", f"{T100['ai']['never_alerted']} 個", f"{T100['rules']['never_alerted']} 個", "—"],
    ["潛在攔阻金額（確認當天凍結）", f"**{wan(T100['ai']['prevented_amount'])} 萬元**", f"{wan(T100['rules']['prevented_amount'])} 萬元", "—"],
], [5.6, 4.4, 3.8, 2.6], bold_first_col=True)
d_ai = T100["ai"]["prevented_by_delay"]
P("**潛在攔阻金額的算法與假設。**覆核確認的人頭帳戶視為當天凍結，之後由被害人直接匯入該帳戶的款項計為潛在攔阻金額；"
  "只計被害人匯入、不計人頭帳戶之間的轉帳，因此同一筆錢經過多層帳戶也不會重複計算。這是仿真情境下的上限：實務上"
  f"調查需要時間，集團也可能改用其他帳戶收款。{T('delay')} 為調查延遲的敏感度分析：延遲 1 天時 AI 仍有 {wan(d_ai['1'])} 萬元，"
  f"延遲 3 天降為 {wan(d_ai['3'])} 萬元，延遲 7 天只剩 {wan(d_ai['7'])} 萬元；冷啟動在導入初期就找出較多帳戶，延遲 7 天"
  f"仍有 {wan(T100['cold']['prevented_by_delay']['7'])} 萬元。提早發現之外，縮短從覆核到處置的時間同樣重要。")
names = [("ai", "AI（只用自家警示資料）"), ("rules", "比較基準：規則"), ("dual", "雙名單（保留 10% 給規則）"),
         ("cold", "冷啟動（參考資料分布相同）"), ("cold_shift", "冷啟動（參考資料分布不同）")]
rows = [["排序方式", "當天凍結", "延遲 1 天", "延遲 3 天", "延遲 7 天"]]
for w, name in names:
    d = T100[w]["prevented_by_delay"]
    rows.append([name] + [wan(d[k]) for k in ("0", "1", "3", "7")])
table("delay", "調查延遲對潛在攔阻金額的影響（萬元；每月覆核 100 個帳戶，6 個月合計）", rows,
      [5.6, 2.7, 2.7, 2.7, 2.7], size=10, align_right=(1, 2, 3, 4), bold_first_col=True)

H(4, "4. AI 的進步從哪裡來、資金流向圖用在哪裡")
m = ev["models"]
ab = ev["ablation"]
P(f"在集團層級交叉驗證下，XGBoost 的 PR-AUC 為 {m['XGBoost（FlowAudit）']['pr_auc']:.3f}（95% 信賴區間 "
  f"{m['XGBoost（FlowAudit）']['pr_auc_ci'][0]:.3f}～{m['XGBoost（FlowAudit）']['pr_auc_ci'][1]:.3f}），比較基準的規則為 "
  f"{m['現行規則計分']['pr_auc']:.3f}；隨機森林（{m['隨機森林']['pr_auc']:.3f}）與 XGBoost 接近，本作品採用 XGBoost 是因為它能"
  f"精確計算 TreeSHAP 解釋且訓練快速。進步從哪裡來？只替 8 條規則的「是否命中」學權重，PR-AUC 仍只有 "
  f"{m['規則加權（依警示資料學權重）']['pr_auc']:.3f}；改用規則背後的 16 項連續數值而不切門檻，就提高到 "
  f"{m['XGBoost（只用 16 項規則指標）']['pr_auc']:.3f}；再加上其餘特徵為 {m['XGBoost（FlowAudit）']['pr_auc']:.3f}。"
  "大部分的進步來自「不把規則切成是或否」，AI 的判斷因此建立在稽核人員熟悉的指標上。")
g = ab["資金網路圖"]
P(f"消融實驗（每次拿掉一類特徵重新訓練；全部特徵為 {ab['全部特徵']['pr_auc']:.3f}）顯示，拿掉 {g['n_features']} 項資金網路圖特徵"
  f"（含集中匯入、分散匯出、資金循環與 PageRank 等）後 PR-AUC 只降到 {g['without']:.3f}，只用這些特徵則為 {g['only']:.3f}："
  "這些資訊大多能被其他特徵取代。因此本作品**不把排序成效歸功於資金流向圖**；資金流向圖的主要用途在調查與舉證——"
  f"切出疑似集團（{F('shot_group')}）、提供時間一致的資金循環證據，並作為規則保底名單的依據。")
gi, tr = ev["graph"]["groups"], ev["graph"]["trace"]["downstream_rule_hit"]
miss = "、".join(f"{k} {v} 個" for k, v in gi["uncovered_by_scheme"].items())
P(f"以仿真資料的答案檢驗這個調查用途：系統切出的 {gi['n_groups']} 個疑似集團中，{gi['single_fraud_group']} 個都只對應單一"
  f"真實詐騙集團，涵蓋 {gi['true_groups']} 個真實集團中的 {gi['true_groups_covered']} 個（未涵蓋的是{miss}，由規則保底名單補上）；"
  f"{gi['never_alerted']} 個從未被警示的人頭帳戶中，{gi['never_alerted_linked']} 個和已警示帳戶被切在同一個集團，調查時可以從已知案件"
  f"連結過去。若從已警示帳戶沿資金往下游追一層、再篩出命中紅旗規則的帳戶，{tr['n']} 個帳戶中有 {tr['mules']} 個是警方尚未通報的"
  f"人頭帳戶（占 {tr['mules'] / tr['n']:.0%}，隨機抽樣約 {ev['graph']['trace']['base_rate']:.1%}），其中 {tr['cycle_mules']} 個是 AI 不易"
  "找到的循環交易。")

H(4, "5. AI 何時會失效與因應")
P(f"本研究主動設計實驗找出 AI 會失效的情況，並據此設計系統的因應機制（{T('dilemmas')}、{F('solutions')}）。")
fb = ev["fallback"]
sy, sy2 = fb["tw_sim"]["system"], fb["tw_sim_seed2"]["system"]
table("dilemmas", "AI 的三個困境與本作品的解法", [
    ["困境", "解法", "結果"],
    ["導入初期警示資料太少：第 1 個月只知道 7 個警示帳戶，AI 只找到 7 個",
     "冷啟動：自家警示帳戶不到 50 個時，先用參考資料訓練的模型",
     f"6 個月找到 {T100['cold']['unique_mules']} 個（AI {T100['ai']['unique_mules']} 個）、平均比警方早 "
     f"{T100['cold']['lead_days_mean']:.0f} 天；參考資料分布不同時仍有 {T100['cold_shift']['unique_mules']} 個"],
    [f"沒見過的詐騙手法：訓練時拿掉一整類手法，AI 只找回該類的 {UNSEEN}",
     "雙名單：保留 10% 覆核名額給規則；每月重新訓練",
     f"平常不減損（逐月模擬 {T100['dual']['unique_mules']} 個，AI 單獨 {T100['ai']['unique_mules']} 個）；新手法累積 "
     f"10 個警示帳戶後，找回率回升到 {LEARN10}"],
    [f"循環交易很少被通報：AI 只找到 {ev['errors']['recall_by_scheme']['循環交易']:.0%}",
     "規則名單＋規則保底名單",
     f"規則名單的 20 個名額中有 {sy['rule_list_cycle']} 個是循環交易帳戶，加上保底名單，"
     f"{sy['rule_list_cycle'] + sy['fallback_cycle']} 個全部找到；另一份資料 {sy2['rule_list_cycle'] + sy2['fallback_cycle']} 個也全部找到"],
], [5.0, 4.9, 6.5])
figure("solutions", A / "fig_solutions.png", 15.5, "各種排序方式的 6 個月成果（逐月模擬，每月覆核 100 個帳戶）")
ls = ev["label_scarcity"]["0.05"]
P(f"此外，當銀行只知道 5%（{ls['n_train_positive']} 個）的警示帳戶時，AI 的 PR-AUC 仍有 {ls['pr_auc']:.3f}，高於規則的 "
  f"{m['現行規則計分']['pr_auc']:.3f}，代表不需要大量警示資料就能起步。")

H(4, "6. 穩健性測試與外部資料驗證")
st = ev["stress"]
ai_re = [v["ai_pr_auc"] for v in st.values()]
ai_tr = [v["transfer_ai_pr_auc"] for k, v in st.items() if k != "基準（新隨機種子）"]
ru = [v["rules_pr_auc"] for v in st.values()]
mt = json.loads((ROOT / "outputs/tw_sim/metrics.json").read_text(encoding="utf-8"))
s2 = mt["external_validation"]["tw_sim_seed2"]
P(f"為了檢查模型是否只是「學會了自己的仿真器」，本研究的驗證分為三類（{F('external')}）：")
numbered([f"**仿真穩健性測試（同一個仿真器）**：以另一個隨機種子產生、模型從未見過的仿真資料直接評分，AI 為 "
          f"{s2['model_pr_auc']:.3f}、規則為 {s2['rules_pr_auc']:.3f}；壓力測試讓人頭帳戶更隱蔽、轉出更慢、收款更分散、"
          f"正常帳戶更像人頭，AI 重新訓練後為 {min(ai_re):.3f}～{max(ai_re):.3f}、不重新訓練為 {min(ai_tr):.3f}～{max(ai_tr):.3f}，"
          f"規則則降到 {min(ru):.3f}～{max(ru):.3f}。這一類只能說明在仿真器能產生的變化內結果穩定。",
          "**獨立來源合成資料的重新訓練驗證（IBM AMLworld）**：AMLworld 是 IBM 以另一套模擬器產生的公開合成資料 [9][15]，"
          "與本研究的仿真器無關。本研究以同一套程式與參數（未修改任何程式碼與規則門檻），在 AMLworld 上**以該資料本身的"
          "標籤重新訓練，並用 5 折交叉驗證評估**。因此它驗證的是「方法」在另一個資料來源上仍然有效，而不是把台灣資料"
          "訓練好的模型直接拿去使用。",
          "**真實銀行資料**：尚待驗證，是實際導入的第一步（見「貳、二」的導入路徑）。"])
figure("external", A / "fig_external.png", 15.5, "仿真穩健性測試與 AMLworld 驗證的 PR-AUC")
at = {r["k"]: r for r in ams["model"]["at_k"]}
n_ml = AMS_SUM["n_labeled_positive"]
check("AMLworld 覆核表最後一列＝洗錢帳戶數", n_ml, max(at))
check("AMLworld 交易筆數（萬）", 167, round(AMS_SUM["n_transactions"] / 1e4))
check("AMLworld 帳戶數（萬）", 16.2, round(AMS_SUM["n_accounts"] / 1e4, 1))
P(f"AMLworld 取美元交易 167 萬筆、16.2 萬個帳戶，洗錢帳戶占 {ams['positive_rate']:.2%}（{n_ml:,} 個）。AI 的 PR-AUC 為 "
  f"{ams['model']['pr_auc']:.3f}，約為隨機的 {ams['model']['pr_auc'] / ams['positive_rate']:.0f} 倍；規則幾乎失效"
  f"（{ams['rules']['pr_auc']:.3f}），因為規則是依台灣人頭帳戶的態樣設計，且該資料沒有現金、裝置、開戶日等欄位。"
  f"對稽核人員更直觀的是「投入多少覆核人力，能找到多少洗錢帳戶」（{T('amlworld')}）：覆核前 200 名時幾乎全是洗錢帳戶"
  f"（{at[200]['hits']} 個），但只占全部洗錢帳戶的 {at[200]['recall']:.1%}；覆核前 {n_ml:,} 名（等於洗錢帳戶總數）時"
  f"找回 {at[n_ml]['recall']:.1%}。成效明顯低於本研究的仿真資料，本研究的仿真資料應視為偏容易；該資料也沒有集團資訊，"
  "無法做集團層級交叉驗證。")
rows = [["覆核帳戶數", "找到的洗錢帳戶", "命中率（找到÷覆核）", f"找回比例（找到÷{n_ml:,}）"]]
for k in (50, 200, 1000, n_ml):
    r = at[k]
    rows.append([f"{k:,}" + ("（＝洗錢帳戶總數）" if k == n_ml else ""), f"{r['hits']:,}", f"{r['precision']:.1%}", f"{r['recall']:.1%}"])
table("amlworld", "IBM AMLworld：依 AI 風險覆核時找到的洗錢帳戶", rows, [4.6, 3.4, 3.8, 4.2], align_right=(1, 2, 3))

# ------------------------------------------------------------------ 二、應用產業與對象
H(2, "二、應用產業與對象")
P(f"FlowAudit 的輸入只需要交易明細（匯出帳戶、匯入帳戶、金額、日期）；若另有時間、管道、登入裝置與開戶資料，即可發揮完整功能，"
  f"因此不限於單一機構或資料格式。{T('apps')} 為主要的應用對象。")
table("apps", "應用對象與價值", [
    ["對象", "使用方式", "價值"],
    ["銀行法遵、洗錢防制單位", "每月全查、依覆核名單調查、作為疑似洗錢申報佐證", "提早發現人頭帳戶，減少誤報與覆核人力"],
    ["銀行內部稽核", "獨立全查，檢驗既有監控機制是否漏掉可疑帳戶", "全查取代抽樣，作為控制有效性測試的佐證"],
    ["電子支付、純網銀、證券期貨", "同一套流程套用於其金流資料", "帳戶多、交易快的機構更需要自動化"],
    ["信用合作社、農漁會信用部", "以離線儀表板與 Excel 工作底稿運作", "不需建置大型系統即可導入"],
    ["會計師事務所", "查核或洗錢防制諮詢時，對客戶資料全查", "提供可追溯的分析報告與工作底稿"],
    ["主管機關", "金融檢查時對受檢機構資料全查", "快速找出高風險帳戶與疑似集團"],
], [4.2, 6.4, 5.8], bold_first_col=True)
cost = T100["reviews"] * 0.5 * 800  # 每件 30 分鐘、時薪 800 元（假設值）
check("覆核帳戶數", 600, T100["reviews"])
P(f"**未來產業價值。**在成本面，以每件覆核 30 分鐘、稽核人員時薪 800 元（假設值）粗估，6 個月覆核 600 個帳戶約需 "
  f"{cost / 1e4:.0f} 萬元人力成本，仿真情境下的潛在攔阻金額約為其 {T100['ai']['prevented_amount'] / cost:.0f} 倍"
  f"（調查延遲 3 天時約 {d_ai['3'] / cost:.0f} 倍）。此比值未計入系統建置、誤報處理與後續調查成本，也不是投資報酬率，"
  "只用來說明量級。在法遵面，規則逐條對應態樣與法規、報告與工作底稿可追溯，有助於落實異常帳戶管理與疑似洗錢申報；"
  "在產業面，冷啟動的「參考資料模型」可發展為跨機構聯防——機構之間分享模型而非原始交易，兼顧個資保護。")
P("**導入路徑。**建議分三階段：（1）以機構的歷史交易與警示紀錄重新訓練與驗證，確認成效與誤報；（2）與現行監控系統平行運作數個月，"
  "比較兩者找到的帳戶；（3）正式納入每月作業，覆核結論回饋模型，並依金管會〈金融業運用人工智慧（AI）指引〉[8] 建立模型治理，"
  "定期檢視成效、記錄版本與參數。系統匯出的 CSV 與 Excel 也可再匯入 JCAATs 等稽核軟體進一步分析。")

# ------------------------------------------------------------------ 三、亮點
H(2, "三、亮點")
P("本作品對台灣解決的問題與提供的便利如下：")
numbered([
    "**聚焦台灣的人頭帳戶問題**：全台仍有約 14.85 萬個警示帳戶 [5]，人頭帳戶是詐騙金流的必經環節。8 條紅旗規則逐條對應"
    "調查局疑似洗錢態樣與金管會管理辦法，仿真資料依台灣的分層轉帳、車手提領、ATM 限額等實況設計，產出的報告與工作底稿"
    "可作為異常帳戶管理與疑似洗錢申報的佐證資料。",
    f"**讓有限的人力用在刀口上**：覆核同樣 {n_pos} 個帳戶，AI 名單中的正常帳戶為 {eq[n_pos]['ai']['normal']:.0f} 個，"
    f"比較基準的規則為 {eq[n_pos]['rules']['normal']:.0f} 個；要找回八成人頭帳戶，AI 需覆核 {r80['ai']['reviewed']:,.0f} 個帳戶，"
    f"規則需覆核 {r80['rules']['reviewed']:,.0f} 個。",
    f"**提早發現、爭取攔阻時間**：逐月模擬中 {T100['ai']['caught_before_alert']} 個人頭帳戶比警方通報平均早 "
    f"{T100['ai']['lead_days_mean']:.0f} 天被找出；潛在攔阻金額在確認當天凍結時為 {wan(d_ai['0'])} 萬元，"
    f"調查延遲 3 天降為 {wan(d_ai['3'])} 萬元，提醒機構同時縮短調查到處置的時間。",
    "**中小型機構也能導入**：完全離線運作，在一台 16 核心的個人電腦上約 1.5 分鐘即可處理 167 萬筆交易；冷啟動讓警示資料很少的機構也能起步；"
    "未來可發展為分享模型、而非分享原始交易的跨機構聯防。",
    "**可信任的 AI**：每個帳戶都附規則證據與 AI 判斷依據；本研究主動測試 AI 會失效的情況並提出因應，清楚區分已證實與"
    "仍屬推估的結果，呼應金管會 AI 指引重視的透明性與可解釋性 [8]。",
])
callout("技術亮點", [
    "**時間一致的資金循環搜尋**：從每一筆交易出發，只沿時間先後一致、金額相近（下一手為上一手的 0.5～1.2 倍）的交易往下找，"
    "5 步內回到起點才算循環，避免把無關的往來誤認為循環。搜尋時加入精確剪枝（只剩 k 步時，下一個帳戶必須能在 k 步內回到起點），"
    "找到的循環與逐一搜尋完全相同，大資料時快約 9 倍。",
    "**防止資料外洩的評估設計**：特徵只用交易與帳戶屬性；只用已警示帳戶訓練；同一詐騙集團的帳戶不會同時出現在訓練與測試。",
    "**公平比較**：AI 與規則在相同覆核量、相同找回率下比較，同分帳戶的覆核順序取期望值，避免比較標準不一致。",
    "**大資料效能與可重現**：IBM AMLworld 167 萬筆交易、16.2 萬個帳戶，全流程約 1.5 分鐘、記憶體峰值 1.6 GB（16 核心個人電腦）；"
    "固定亂數種子、鎖定套件版本，並以 20 項自動化測試確保每次修改後結果一致。",
])

# ================================================================== 參、結論
H(1, "參、結論")
H(2, "一、結論")
P(f"FlowAudit 以「全查、看資金流向、規則與 AI 分工、持續稽核」的方式偵測人頭帳戶。在台灣情境仿真資料上，每月只覆核 100 個"
  f"帳戶，6 個月即找出 {T100['ai']['unique_mules']} 個人頭帳戶，其中 {T100['ai']['caught_before_alert']} 個比警方通報平均早 "
  f"{T100['ai']['lead_days_mean']:.0f} 天；覆核同樣多的帳戶時，名單中的正常帳戶明顯少於比較基準的規則。本作品也主動找出 AI "
  "會失效的情況，以冷啟動、三份名單與每月重新訓練因應，並在獨立來源的合成資料 AMLworld 上以其標籤重新訓練，確認方法仍然有效。")
H(2, "二、限制")
numbered([
    "主要成果來自本研究自行產生的合成資料。仿真器依公開資料設計，並刻意加入大量像人頭帳戶的正常帳戶，但仍不等於真實銀行資料；"
    "真實資料上的成效尚待驗證。",
    f"AMLworld 也是合成資料，且是以其標籤重新訓練的結果，只能說明方法可用於其他資料來源，不能說明已訓練的模型可以直接跨資料使用；"
    f"其成效（PR-AUC {ams['model']['pr_auc']:.3f}、覆核 {n_ml:,} 名時找回 {at[n_ml]['recall']:.0%}）遠低於本研究的仿真資料，"
    "也沒有集團資訊可做集團層級交叉驗證。",
    f"潛在攔阻金額假設覆核確認後即可凍結、集團不會改用其他帳戶，應視為上限；調查延遲 3 天時即降為 {wan(d_ai['3'])} 萬元。",
    "冷啟動的參考資料與自家資料出自同一個仿真器，效果應視為上限；實際跨機構資料的差異會更大。",
    "規則門檻依台灣人頭帳戶的態樣設計，應用於不同機構或資料時需重新設定；資金網路圖特徵對排序的額外貢獻很小，其價值主要在調查與舉證。",
])
H(2, "三、未來發展")
numbered([
    "與金融機構合作，以去識別化的真實資料驗證並校正規則門檻與模型。",
    "加入發票、營業額等外部資料，強化循環交易（例如虛增營收）的辨識。",
    "嘗試圖神經網路等能直接學習資金網路結構的模型，並與現行方法比較。",
    "發展跨機構聯防：在不交換原始交易的前提下分享模型或已知態樣。",
    "依金管會 AI 指引建立模型治理與監控機制，定期檢視模型是否失準。",
])

# ================================================================== 肆、參考資料
H(1, "肆、參考資料")
refs = [
    "法務部調查局，〈疑似洗錢或資恐交易態樣〉（銀行業）。",
    "金融監督管理委員會，《存款帳戶及其疑似不法或顯屬異常交易管理辦法》。",
    "金融監督管理委員會，《金融機構防制洗錢辦法》。",
    "台視新聞（Yahoo 新聞轉載），〈2025年全台詐騙財損近900億！165打詐儀錶板揭最常見手法〉，2026 年 1 月 19 日。",
    "NOWnews，〈「警示帳戶」連3降！1月再減2232戶 但全台還有14.85萬戶被鎖住〉，2026 年 2 月 26 日。",
    "報導者，〈人頭帳戶氾濫，打詐國家隊怎麼加快部會和企業合作、從源頭「阻詐」？〉，2023 年 3 月 15 日。",
    "聯合新聞網，〈洗車場暗藏詐團金流鏈…「五層式」企業分工 屏檢溯源瓦解〉，2026 年 6 月 12 日。",
    "金融監督管理委員會，〈金融業運用人工智慧（AI）指引〉，2024 年。",
    "Altman, E., Blanuša, J., von Niederhäusern, L., Egressy, B., Anghel, A., & Atasu, K. (2023). Realistic Synthetic Financial "
    "Transactions for Anti-Money Laundering Models. NeurIPS 2023 Datasets and Benchmarks Track.",
    "IBM. AMLSim: Anti-Money Laundering Simulator. https://github.com/IBM/AMLSim",
    "Chen, T., & Guestrin, C. (2016). XGBoost: A Scalable Tree Boosting System. Proceedings of KDD 2016, 785–794.",
    "Lundberg, S. M., et al. (2020). From local explanations to global understanding with explainable AI for trees. "
    "Nature Machine Intelligence, 2, 56–67.",
    "Blondel, V. D., Guillaume, J.-L., Lambiotte, R., & Lefebvre, E. (2008). Fast unfolding of communities in large networks. "
    "Journal of Statistical Mechanics: Theory and Experiment, P10008.",
    "Money101，〈全台30家銀行ATM跨行提款、轉帳金額上限一覽表〉。",
    "IBM Transactions for Anti Money Laundering (AML)，Kaggle 資料集（本研究使用 HI-Small_Trans.csv）。"
    "https://www.kaggle.com/datasets/ealtman2019/ibm-transactions-for-anti-money-laundering-aml",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(-30)
    p.paragraph_format.left_indent = Pt(30)
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    runs(p, f"[{i}] {r}")  # 參考資料也是內文，用 12 點

# ------------------------------------------------------------------ 核對與存檔
assert FIG[0] == len(FIGS) and TAB[0] == len(TABS), (FIG[0], TAB[0])
check("逐月模擬 AI 找到", 97, T100["ai"]["unique_mules"])
check("逐月模擬 規則找到", 62, T100["rules"]["unique_mules"])
check("逐月模擬 被害人匯入總額（萬元）", 9420, round(ev["temporal"]["summary"]["victim_amount_total"] / 1e4))
check("第 1 個月已知警示帳戶", 7, ev["temporal"]["batches"][0]["n_known_alerts"])
check("第 1 個月 AI 找到", 7, ev["temporal"]["batches"][0]["by_k"]["100"]["ai"]["tp"])
check("冷啟動切換門檻", 50, ev["temporal"]["settings"]["switch_after"])
bad =[(lab, a, b) for lab, a, b in CHECKS if a != b]
if bad:
    raise SystemExit(f"正文數字與結果檔不一致：{bad}")

cp = doc.core_properties
cp.author = cp.last_modified_by = cp.comments = ""
cp.title = "FlowAudit 金流偵探 企劃書"
doc.save(OUT)
print("saved", OUT, "| 圖", FIG[0], "| 表", TAB[0], "| 核對", len(CHECKS), "項")
